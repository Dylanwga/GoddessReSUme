"""Build editable resume blocks by cloning the bundled, anonymized DOCX.

Requires python-docx and lxml (use the bundled Python runtime). Coordinates and
sizes are points; pages are one-based. The eight roles are name, section,
project, body, bullet, citation, contact, and institution. Every call creates
one native textbox containing one logical paragraph. A newline in supplied
text is an explicit line break, not another paragraph.

This module does NOT measure text, detect visual overlap, or paginate. Allocate
enough height and render every page with the bundled LibreOffice renderer after
saving; passing validate() establishes structural bounds, not visual fitness.
Only caller-supplied text, links, and pictures are inserted. template_path must
be a compatible anonymized component template, never a personal source resume.

Example::

    b = ResumeBuilder(theme="2F6FBA")
    b.add_block("name", "虚构候选人", y=28, width=250, height=36)
    b.add_block("section", "虚构项目", y=110)
    b.add_block("bullet", [{"text": "职责：", "bold": True},
                          {"text": "调用者提供的内容。"}], y=150, height=42)
    b.save("fictional.docx")  # Then render and inspect; no auto-fit is implied.
"""

from copy import deepcopy
from dataclasses import dataclass
import math
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "layout-reference.docx"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
PROTOTYPES = {
    "name": "姓名占位", "section": "教育背景",
    "project": "项目 A：训练工具与实验验证", "body": "背景：",
    "bullet": "技术实现：", "citation": "虚构论文题录",
    "contact": "代码主页占位", "institution": "示例大学甲",
}
SIZES = {"name": 20, "section": 14, "contact": 11, "institution": 11}
HEIGHTS = {"name": 36, "section": 28, "contact": 24, "institution": 28,
           "project": 24, "body": 40, "bullet": 48, "citation": 48}


def _child(parent, tag, **attributes):
    node = parent.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        parent.append(node)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    return node


def _hex(value):
    value = str(value).lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", value):
        raise ValueError(f"Expected six-digit RGB color, got {value!r}")
    return value


def _number(value, name, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{name} must be a finite number of points")
    if value < 0 or (positive and value == 0):
        raise ValueError(f"{name} must be {'positive' if positive else 'nonnegative'}")
    return value


@dataclass
class Block:
    """A caller-addressable block; paragraph supports normal python-docx edits."""

    role: str
    page: int
    anchor: object
    paragraph: Paragraph


class ResumeBuilder:
    """Reuse native components, numbering, and styles from the safe template.

    theme is a caller-selected RGB color; omission uses neutral gray 4A5568.
    This helper does not infer colors from logos. Font defaults reproduce the
    reference and must exist when rendering.
    The original template is read only, and save() refuses to overwrite it.
    """

    def __init__(self, template_path=None, *, theme=None,
                 east_asia_font="KaiTi", western_font="Times New Roman"):
        self.template_path = Path(template_path or TEMPLATE).resolve()
        self.document = Document(self.template_path)
        self.theme = _hex(theme if theme is not None else "4A5568")
        self.dark = "".join(f"{round(int(self.theme[i:i + 2], 16) * .8):02X}" for i in (0, 2, 4))
        self.east_asia_font, self.western_font = east_asia_font, western_font
        self._prototypes, self._pages, self.blocks, self._next_id = {}, {}, [], 1
        for anchor in self.document.element.body.iter(qn("wp:anchor")):
            prop = anchor.find(qn("wp:docPr"))
            if prop is not None:
                self._prototypes.setdefault(prop.get("name"), deepcopy(anchor))
        missing = set(PROTOTYPES.values()) - self._prototypes.keys()
        if missing:
            raise ValueError(f"Template lacks named anonymized components: {sorted(missing)}")
        body = self.document.element.body
        section = deepcopy(body.sectPr)
        for node in list(section):
            if node.tag in {qn("w:headerReference"), qn("w:footerReference")}:
                section.remove(node)
        body.clear()
        body.append(section)
        keep = {"styles", "stylesWithEffects", "settings", "webSettings", "fontTable", "theme", "numbering"}
        for rid, rel in list(self.document.part.rels.items()):
            if rel.is_external or rel.reltype.rsplit("/", 1)[-1] not in keep:
                del self.document.part.rels[rid]
        for field in ("author", "last_modified_by", "title", "subject", "keywords", "comments"):
            setattr(self.document.core_properties, field, "")

    def _id(self):
        value = self._next_id
        self._next_id += 1
        return value

    def _page(self, page):
        if isinstance(page, bool) or not isinstance(page, int) or page < 1:
            raise ValueError("page must be a positive, one-based integer")
        for number in range(len(self._pages) + 1, page + 1):
            host = self.document.add_paragraph()
            props = host._p.get_or_add_pPr()
            _child(props, "w:spacing", before=0, after=0, line=20, lineRule="exact")
            if number > 1:
                _child(props, "w:pageBreakBefore")
            self._pages[number] = host
        return self._pages[page]

    def _font(self, run, size, color, bold=False, italic=False, underline=False):
        run.font.size, run.font.color.rgb = Pt(size), RGBColor.from_string(_hex(color))
        run.bold, run.italic, run.underline = bold, italic, underline
        fonts = _child(run._r.get_or_add_rPr(), "w:rFonts")
        fonts.attrib.clear()
        for key in ("ascii", "hAnsi", "cs"):
            fonts.set(qn("w:" + key), self.western_font)
        fonts.set(qn("w:eastAsia"), self.east_asia_font)
        _child(run._r.get_or_add_rPr(), "w:szCs", val=round(size * 2))

    def _picture(self, paragraph, path, height, size, *, baseline=0):
        _number(height, "icon_height", positive=True)
        run = paragraph.add_run()
        self._font(run, size, "262626")
        shape = run.add_picture(str(path), height=Pt(height))
        identifier = self._id()
        for node in shape._inline.iter():
            if node.tag in {qn("wp:docPr"), qn("pic:cNvPr")}:
                node.set("id", str(identifier))
                node.set("name", f"image-{identifier}")
        if baseline:
            _child(run._r.get_or_add_rPr(), "w:position", val=round(baseline * 2))
        return run

    def add_block(self, role, runs_or_text, *, page=1, x=34, y=28,
                  width=None, height=None, font_size=None, color=None, level=1,
                  icon_path=None, icon_height=11):
        """Add one paragraph in one textbox; positions/sizes are caller-owned.

        runs_or_text is a string, a run dict, or a sequence of either. Run keys:
        text, bold, italic, underline, color, url, font_size. An icon-only run
        accepts icon_path and optional icon_height (default 11). body is plain; bullet uses
        native numbering at level 0 (solid) or 1 (hollow). Section rules remain
        2 pt. Icons are inline and embedded in the document part, not linked.
        """
        if role not in PROTOTYPES:
            raise ValueError(f"Unknown role {role!r}; choose {tuple(PROTOTYPES)}")
        if level not in (0, 1):
            raise ValueError("bullet level must be 0 or 1")
        section = self.document.sections[0]
        width = section.page_width.pt - x - 34 if width is None else width
        height = HEIGHTS[role] if height is None else height
        size = SIZES.get(role, 10.5) if font_size is None else font_size
        for value, name in ((x, "x"), (y, "y"), (width, "width"), (height, "height"), (size, "font_size")):
            _number(value, name, positive=name in {"width", "height", "font_size"})
        pieces = [runs_or_text] if isinstance(runs_or_text, (str, dict)) else list(runs_or_text)
        pieces = [{"text": item} if isinstance(item, str) else dict(item) for item in pieces]
        allowed = {"text", "bold", "italic", "underline", "color", "url", "font_size",
                   "icon_path", "icon_height"}
        for item in pieces:
            if set(item) - allowed or not isinstance(item.get("text", ""), str):
                raise ValueError(f"Invalid run specification: {item!r}")
            if "color" in item:
                _hex(item["color"])
            if item.get("url") and not isinstance(item["url"], str):
                raise ValueError("url must be a string")
            _number(item.get("font_size", size), "run font_size", positive=True)
            if "icon_path" in item:
                if item.get("text") or item.get("url"):
                    raise ValueError("An icon run cannot also contain text or a URL")
                _number(item.get("icon_height", 11), "icon_height", positive=True)
        default_color = color or (self.theme if role == "section" else self.dark
                                 if role in {"project", "institution"} else "262626")
        default_color = _hex(default_color)
        anchor = deepcopy(self._prototypes[PROTOTYPES[role]])
        data = next(anchor.iter(qn("a:graphicData")))
        shape = deepcopy(next(anchor.iter("{" + WPS + "}wsp")))
        data.clear()
        data.set("uri", WPS)
        data.append(shape)  # The institution's sample badge is deliberately discarded.
        extent = {"cx": str(round(width * 12700)), "cy": str(round(height * 12700))}
        anchor.find(qn("wp:extent")).attrib.update(extent)
        next(shape.iter(qn("a:ext"))).attrib.update(extent)
        for axis, value in (("H", x), ("V", y)):
            position = anchor.find(qn("wp:position" + axis))
            position.set("relativeFrom", "page")
            position.find(qn("wp:posOffset")).text = str(round(value * 12700))
        identifier = self._id()
        prop = anchor.find(qn("wp:docPr"))
        prop.attrib.update({"id": str(identifier), "name": f"{role}-{identifier}", "descr": ""})
        anchor.set("relativeHeight", str(251660000 + identifier))
        tx = next(shape.iter(qn("w:txbxContent")))
        paragraph_xml = deepcopy(next(tx.iter(qn("w:p"))))
        for child in list(paragraph_xml):
            if child.tag != qn("w:pPr"):
                paragraph_xml.remove(child)
        tx.clear()
        tx.append(paragraph_xml)
        paragraph = Paragraph(paragraph_xml, self.document._body)
        props = paragraph_xml.get_or_add_pPr()
        if role == "body":
            for tag in ("w:numPr", "w:ind"):
                node = props.find(qn(tag))
                if node is not None:
                    props.remove(node)
        if role == "bullet":
            _child(props.find(qn("w:numPr")), "w:ilvl", val=level)
            _child(props, "w:ind", left=142 if level == 0 else 255, hanging=142)
        if role == "section":
            _child(_child(props, "w:pBdr"), "w:bottom", val="single", sz=16, space=2, color=self.theme)
        if role == "institution":
            next(shape.iter(qn("a:srgbClr"))).set("val", self.theme)
            _child(props, "w:ind", left=0, hanging=0)
        if icon_path is not None:
            self._picture(paragraph, icon_path, icon_height, size, baseline=-1 if role == "contact" else 0)
            self._font(paragraph.add_run(" "), size, default_color)
        for item in pieces:
            run_size = item.get("font_size", size)
            if "icon_path" in item:
                self._picture(paragraph, item["icon_path"], item.get("icon_height", 11),
                              run_size, baseline=-1 if role == "contact" else 0)
                continue
            run = paragraph.add_run(item.get("text", ""))
            self._font(run, run_size, item.get("color", "0000FF" if item.get("url") else default_color),
                       item.get("bold", role in {"name", "section", "project", "institution", "citation"}),
                       item.get("italic", role == "citation"), item.get("underline", bool(item.get("url"))))
            if item.get("url"):
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), self.document.part.relate_to(item["url"], RT.HYPERLINK, is_external=True))
                link.append(run._r)
                paragraph_xml.append(link)
        drawing = OxmlElement("w:drawing")
        drawing.append(anchor)
        self._page(page).add_run()._r.append(drawing)
        block = Block(role, page, anchor, paragraph)
        self.blocks.append(block)
        return block

    def add_institution(self, name, detail="", degree="", dates="", *,
                        logo_path=None, tab_stops=None, **placement):
        """One tinted textbox with inline logo and three real tab stops.

        tab_stops is three point offsets from the textbox left edge. The last
        stop is right-aligned. Long fields must be inspected after rendering.
        """
        fields = (name, detail, degree, dates)
        if any(not isinstance(value, str) or any(c in value for c in "\t\r\n") for value in fields):
            raise ValueError("Institution fields must be single-line strings without tabs")
        width = placement.get("width")
        if width is None:
            width = self.document.sections[0].page_width.pt - placement.get("x", 34) - 34
        _number(width, "width", positive=True)
        stops = tuple(tab_stops) if tab_stops is not None else (width * .37, width * .64, width - 3)
        for stop in stops:
            _number(stop, "tab stop", positive=True)
        if len(stops) != 3 or any(not 0 < stop <= width for stop in stops) or list(stops) != sorted(set(stops)):
            raise ValueError("tab_stops must be three increasing positive offsets within the textbox")
        block = self.add_block("institution", "\t".join(fields), icon_path=logo_path,
                               icon_height=18, **placement)
        tabs = _child(block.paragraph._p.get_or_add_pPr(), "w:tabs")
        tabs.clear()
        for index, position in enumerate(stops):
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "right" if index == 2 else "left")
            tab.set(qn("w:pos"), str(round(position * 20)))
            tabs.append(tab)
        return block

    def add_contact(self, text, *, url=None, icon_path=None, **placement):
        """Embed an optional caller-provided icon; link only to the supplied URL."""
        return self.add_block("contact", {"text": text, "url": url},
                              icon_path=icon_path, **placement)

    def validate(self):
        """Raise on invalid geometry/structure/relationships; never trim content."""
        section = self.document.sections[0]
        page_width, page_height = section.page_width.pt, section.page_height.pt
        seen = set()
        for block in self.blocks:
            anchor = block.anchor
            extent = anchor.find(qn("wp:extent"))
            width, height = (int(extent.get(k)) / 12700 for k in ("cx", "cy"))
            x, y = (int(anchor.find(qn("wp:position" + a)).find(qn("wp:posOffset")).text) / 12700 for a in ("H", "V"))
            if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > page_width + 1e-5 or y + height > page_height + 1e-5:
                raise ValueError(f"{block.role} on page {block.page} exceeds page bounds: {(x, y, width, height)}")
            boxes = list(anchor.iter(qn("w:txbxContent")))
            if len(boxes) != 1 or len(list(boxes[0].iter(qn("w:p")))) != 1:
                raise ValueError(f"{block.role} must have exactly one textbox and one paragraph")
        for node in self.document.element.body.iter():
            if node.tag == qn("wp:docPr"):
                identifier = node.get("id")
                if identifier in seen:
                    raise ValueError(f"Duplicate drawing id {identifier}")
                seen.add(identifier)
            for key in ("r:id", "r:embed", "r:link"):
                rid = node.get(qn(key))
                if rid is not None and rid not in self.document.part.rels:
                    raise ValueError(f"Unresolved document relationship {rid}")
        return {"pages": len(self._pages), "textboxes": len(self.blocks),
                "requires_rendering": True, "measures_text_fit": False}

    def save(self, path):
        """Validate, then save a new DOCX. Render it before final delivery."""
        path = Path(path).resolve()
        if path == self.template_path:
            raise ValueError("The component template cannot be overwritten")
        self.validate()
        self.document.save(path)
        return path
